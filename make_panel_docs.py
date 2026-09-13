"""Generate the panel deliverables from the markdown sources.

  venv\\Scripts\\python.exe make_panel_docs.py

Produces:
  docs/CBODS_PRESENTATION.html   self-contained browser deck (Inter embedded,
                                 speaker notes, keyboard nav, print/PDF export)
  docs/CBODS_PRESENTATION.pptx   10 editable slides + speaker notes
  docs/PANEL_PREP.pdf            styled A4 handout

The markdown files remain the single source of truth: edit
docs/PRESENTATION_OUTLINE.md or docs/PANEL_PREP.md, re-run this, and all three
follow. Requires python-pptx, fpdf2 (pip install).
"""
import base64
import json
import pathlib
import re
import sys

ROOT = pathlib.Path(__file__).resolve().parent
DOCS = ROOT / "docs"

EMPHASIS = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def strip_md(text):
    return text.replace("**", "").replace("*", "").replace("`", "")


def md_inline(text):
    """Markdown emphasis to inline HTML, with escaping."""
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"\*([^*]+)\*", r"<em>\1</em>", text)
    return text


def parse_outline():
    md = (DOCS / "PRESENTATION_OUTLINE.md").read_text(encoding="utf-8")
    heading = re.compile(r"^## Slide (\d+) — (.+?) \((\d+:\d+)\)\s*$", re.M)
    parts = heading.split(md)
    if len(parts) < 3:
        sys.exit("Could not find slide headings in PRESENTATION_OUTLINE.md")
    slides = []
    for i in range(1, len(parts) - 3, 4):
        number, title, seconds, body = parts[i], parts[i + 1], parts[i + 2], parts[i + 3]
        body = body.split("\n## Delivery tips")[0]
        on_slide, notes_parts, notes, mode = [], [], "", None
        for line in body.splitlines():
            stripped = line.strip()
            if stripped.startswith("**On slide:**"):
                mode = "bullets"
                continue
            if stripped.startswith("**Notes:**"):
                mode = "notes"
                notes_parts = [stripped[len("**Notes:**"):].strip()]
                continue
            if stripped == "---" or not stripped:
                if mode == "notes" and notes_parts:
                    notes = " ".join(notes_parts)
                continue
            if mode == "bullets" and stripped.startswith("- "):
                on_slide.append(stripped[2:])
            elif mode == "bullets" and on_slide:
                # Soft-wrapped continuation of the previous bullet.
                on_slide[-1] += " " + stripped
            elif mode == "notes":
                notes_parts.append(stripped)
        slides.append({"number": int(number), "title": title, "time": seconds, "bullets": on_slide, "notes": notes})
    return slides


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def build_pptx(slides):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    def md_runs(paragraph, text, size, color):
        for token in EMPHASIS.split(text):
            if not token:
                continue
            run = paragraph.add_run()
            run.text = strip_md(token)
            font = run.font
            font.size = Pt(size)
            font.color.rgb = color
            if token.startswith("**"):
                font.bold = True
            elif token.startswith("*"):
                font.italic = True

    BRAND = RGBColor(0xA4, 0x16, 0x1A)
    DARK = RGBColor(0x22, 0x22, 0x22)
    GREY = RGBColor(0x6B, 0x6B, 0x6B)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    LIGHT = RGBColor(0xF3, 0xDE, 0xE1)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_box(slide, x, y, w, h):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.text_frame.word_wrap = True
        return box.text_frame

    def add_rule(slide, x, y, w, color, thick=0.035):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(thick))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        return shape

    total = len(slides)
    for data in slides:
        slide = prs.slides.add_slide(blank)
        if data["number"] == 1:
            background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
            background.fill.solid()
            background.fill.fore_color.rgb = BRAND
            background.line.fill.background()

            tf = add_box(slide, 0.9, 1.7, 11.5, 1.6)
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "CBODS"
            run.font.size = Pt(66)
            run.font.bold = True
            run.font.color.rgb = WHITE

            tf = add_box(slide, 0.95, 3.0, 11.5, 0.7)
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "Centralised Blood & Organ Donation System"
            run.font.size = Pt(26)
            run.font.color.rgb = LIGHT

            tagline, meta = cover_meta(slides)
            tf = add_box(slide, 0.95, 4.3, 11.5, 0.6)
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = tagline
            run.font.size = Pt(18)
            run.font.italic = True
            run.font.color.rgb = LIGHT

            tf = add_box(slide, 0.95, 5.6, 11.5, 1.2)
            for i, line in enumerate(meta):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                run = p.add_run()
                run.text = line
                run.font.size = Pt(16)
                run.font.color.rgb = WHITE
        else:
            tf = add_box(slide, 0.7, 0.45, 10.4, 1.0)
            p = tf.paragraphs[0]
            md_runs(p, data["title"], 30, BRAND)
            for run in p.runs:
                run.font.bold = True

            chip = add_box(slide, 11.4, 0.55, 1.2, 0.4)
            p = chip.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = data["time"]
            run.font.size = Pt(14)
            run.font.color.rgb = GREY

            add_rule(slide, 0.72, 1.45, 11.9, BRAND)

            tf = add_box(slide, 0.85, 1.85, 11.6, 4.9)
            for i, bullet in enumerate(data["bullets"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(14)
                marker = p.add_run()
                marker.text = "•  "
                marker.font.size = Pt(20)
                marker.font.color.rgb = BRAND
                marker.font.bold = True
                md_runs(p, bullet, 20, DARK)

            footer = add_box(slide, 0.7, 7.05, 8.0, 0.35)
            run = footer.paragraphs[0].add_run()
            run.text = "CBODS — Centralised Blood & Organ Donation System"
            run.font.size = Pt(10)
            run.font.color.rgb = GREY

            pager = add_box(slide, 11.5, 7.05, 1.1, 0.35)
            p = pager.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = f"{data['number']} / {total}"
            run.font.size = Pt(10)
            run.font.color.rgb = GREY

        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = data["notes"] or data["title"]

    out = DOCS / "CBODS_PRESENTATION.pptx"
    prs.save(out)
    return out, total


# ---------------------------------------------------------------------------
# HTML deck
# ---------------------------------------------------------------------------

HTML_TEMPLATE = r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>CBODS — Panel Presentation</title>
<style>
/* CBODS design tokens, lifted from static/css/cbods.css so the deck and the
   app read as one product. Inter is embedded, so this file works offline. */
@font-face {
  font-family: "InterVariable";
  font-style: normal;
  font-weight: 100 900;
  font-display: swap;
  src: url(data:font/woff2;base64,__INTER_B64__) format("woff2");
}
:root {
  --brand: #a4161a;
  --on-brand: #fff;
  --bg: #f6f2f3;
  --surface: #ffffff;
  --ink: #1f1a1b;
  --muted: #6b6b6b;
  --hairline: rgba(16, 24, 40, .12);
  --shadow: 0 2px 4px rgba(16, 24, 40, .05), 0 8px 24px rgba(16, 24, 40, .08);
  --radius: 1rem;
}
* { box-sizing: border-box; }
html, body { height: 100%; }
body {
  margin: 0;
  font-family: "InterVariable", system-ui, -apple-system, "Segoe UI", Roboto, sans-serif;
  background: var(--bg);
  color: var(--ink);
  line-height: 1.55;
  overflow: hidden;
}
.deck { height: 100%; display: grid; place-items: center; padding: 1rem 1rem 3.2rem; }
.slide {
  width: min(100%, calc((100vh - 5.2rem) * 16 / 9));
  aspect-ratio: 16 / 9;
  background: var(--surface);
  border-radius: var(--radius);
  box-shadow: var(--shadow);
  padding: clamp(1.2rem, 3.2vmin, 3rem) clamp(1.5rem, 4.2vmin, 4rem);
  display: none;
  flex-direction: column;
}
.slide.active { display: flex; }
.slide-title {
  font-weight: 600;
  letter-spacing: -.014em;
  color: var(--brand);
  font-size: clamp(1.35rem, 4.4vmin, 2.6rem);
  line-height: 1.15;
  margin: 0;
}
.rule { height: 3px; background: var(--brand); border: 0; width: 100%; margin: .7rem 0 0; }
.bullets { margin: clamp(.8rem, 2.4vmin, 2rem) 0 0; padding: 0; list-style: none; flex: 1; min-height: 0; }
.bullets li {
  display: flex;
  gap: .7em;
  font-size: clamp(.98rem, 2.75vmin, 1.7rem);
  margin-bottom: clamp(.4rem, 1.3vmin, 1rem);
}
.bullets .marker { color: var(--brand); font-weight: 700; }
.chip {
  position: absolute;
  top: clamp(1rem, 2.6vmin, 2.2rem);
  right: clamp(1.4rem, 3.6vmin, 3.4rem);
  font-variant-numeric: tabular-nums;
  color: var(--muted);
  font-size: clamp(.75rem, 1.8vmin, 1.05rem);
}
.slide { position: relative; }
.slide-foot {
  display: flex;
  justify-content: space-between;
  color: var(--muted);
  font-size: clamp(.65rem, 1.6vmin, .9rem);
  margin-top: .6rem;
}
.slide.cover {
  background: linear-gradient(135deg, #a4161a 0%, #7d1013 100%);
  color: var(--on-brand);
  justify-content: center;
}
.slide.cover .slide-title { color: var(--on-brand); font-size: clamp(2.6rem, 10vmin, 6rem); letter-spacing: -.03em; }
.slide.cover .sub { font-size: clamp(1.1rem, 3.4vmin, 2.2rem); font-weight: 300; margin: .3rem 0 0; }
.slide.cover .rule { background: rgba(255,255,255,.45); }
.slide.cover .tagline { font-style: italic; opacity: .93; font-size: clamp(.95rem, 2.6vmin, 1.6rem); margin-top: clamp(1rem, 4vmin, 2.4rem); }
.slide.cover .meta { margin-top: clamp(1rem, 4.5vmin, 2.8rem); font-size: clamp(.85rem, 2vmin, 1.25rem); opacity: .88; white-space: pre-line; }
/* On narrow windows the 16:9 card is too short for the content; let it grow
   so nothing clips. Real presentation displays keep the 16:9 stage. */
@media (max-width: 720px) {
  .slide { aspect-ratio: auto; width: 100%; }
  .bullets li { font-size: 1rem; }
}
.bar {
  position: fixed;
  left: 0; right: 0; bottom: 0;
  display: flex;
  align-items: center;
  gap: .5rem;
  padding: .55rem 1rem;
  background: rgba(31, 26, 27, .93);
  color: #eee;
  font-size: .85rem;
  z-index: 10;
}
.bar button {
  font: inherit;
  background: transparent;
  color: inherit;
  border: 1px solid rgba(255,255,255,.35);
  border-radius: .5rem;
  padding: .18rem .7rem;
  cursor: pointer;
}
.bar button:hover, .bar button.on { background: rgba(255,255,255,.16); }
.bar .progress { flex: 1; height: 4px; background: rgba(255,255,255,.18); border-radius: 2px; overflow: hidden; }
.bar .progress i { display: block; height: 100%; background: #fff; width: 0; transition: width 160ms ease; }
.bar .count { font-variant-numeric: tabular-nums; }
.notes-toast {
  position: fixed;
  left: 50%;
  bottom: 3.4rem;
  transform: translateX(-50%);
  max-width: min(90vw, 58rem);
  background: rgba(31, 26, 27, .96);
  color: #fff;
  padding: .8rem 1.2rem;
  border-radius: .8rem;
  font-size: .95rem;
  line-height: 1.5;
  box-shadow: var(--shadow);
  display: none;
  z-index: 11;
}
.notes-toast.show { display: block; }
/* Presenter view (P): the real deck shrinks to the left column; the right
   column holds the next-slide preview, the notes, and a talk timer. */
#presenter-aside { display: none; }
body.presenter #presenter-aside {
  display: flex;
  position: fixed;
  top: 1rem; right: 1rem; bottom: 4.2rem;
  width: min(30rem, 36vw);
  flex-direction: column;
  gap: .9rem;
  z-index: 6;
}
body.presenter .deck {
  position: fixed;
  top: 1rem; left: 1rem; bottom: 4.2rem;
  width: calc(100vw - min(30rem, 36vw) - 3rem);
  height: auto;
  padding: 0;
}
.pv-card { background: var(--surface); border-radius: var(--radius); box-shadow: var(--shadow); }
.pv-label { font-size: .72rem; text-transform: uppercase; letter-spacing: .08em; color: var(--muted); font-weight: 600; margin: 0 0 .35rem; }
#pv-next-stage { aspect-ratio: 16 / 9; position: relative; overflow: hidden; }
#pv-next-holder { position: absolute; top: 0; left: 0; }
#pv-next-holder .slide { display: flex !important; visibility: visible !important; opacity: 1 !important; position: absolute; top: 0; left: 0; margin: 0; box-shadow: none; }
#pv-notes { flex: 1; min-height: 0; overflow-y: auto; padding: 1rem 1.2rem; }
#pv-notes h3 { margin: 0 0 .5rem; color: var(--brand); font-size: 1.05rem; font-weight: 600; letter-spacing: -.01em; }
#pv-notes .pv-body { font-size: 1.02rem; line-height: 1.6; white-space: pre-line; }
.pv-meta { display: flex; align-items: center; gap: .8rem; padding: .7rem 1.2rem; color: var(--muted); font-size: .9rem; }
#pv-timer { font-variant-numeric: tabular-nums; font-weight: 600; color: var(--ink); }
.pv-end { position: absolute; inset: 0; display: grid; place-items: center; color: var(--muted); font-style: italic; background: var(--surface); border-radius: var(--radius); }
@media (max-width: 900px) {
  body.presenter .deck { width: 100vw; bottom: 40vh; }
  /* The deck is fixed, so the aside's static position is the body top; pin
     it explicitly below the deck instead. */
  body.presenter #presenter-aside {
    top: calc(60vh + .6rem);
    bottom: 4.2rem;
    height: auto;
    width: auto;
    left: 1rem;
    right: 1rem;
  }
  body.presenter #pv-next-wrap { display: none; }
}
@media print {
  body { overflow: visible; background: #fff; }
  .deck { display: block; padding: 0; height: auto; }
  .slide { display: flex; position: static; page-break-after: always; margin: 0 auto 1rem; width: 100%; max-width: none; box-shadow: none; border: 1px solid var(--hairline); }
  .bar, .notes-toast, #presenter-aside { display: none; }
}
</style>
</head>
<body>
<main class="deck" id="deck">__SLIDES__</main>
<nav class="bar" aria-label="Slide controls">
  <button id="prev" aria-label="Previous slide">&#8249; Prev</button>
  <button id="next" aria-label="Next slide">Next &#8250;</button>
  <button id="notes" aria-label="Toggle speaker notes">N &middot; Notes</button>
  <button id="present" aria-label="Toggle presenter view">P &middot; Present</button>
  <button id="help" aria-label="Show keyboard help">? &middot; Help</button>
  <div class="progress" aria-hidden="true"><i id="progress-bar"></i></div>
  <span class="count"><span id="cur">1</span> / <span id="total">__TOTAL__</span></span>
</nav>
<div class="notes-toast" id="toast" role="status"></div>
<aside id="presenter-aside" aria-label="Presenter view">
  <div id="pv-next-wrap">
    <p class="pv-label">Next</p>
    <div class="pv-card" id="pv-next-stage"><div id="pv-next-holder"></div></div>
  </div>
  <div class="pv-card" id="pv-notes">
    <h3 id="pv-title">Slide 1</h3>
    <div class="pv-body" id="pv-body"></div>
  </div>
  <div class="pv-card pv-meta">
    <span id="pv-timer">00:00</span>
    <span id="pv-count">1 / 10</span>
    <span style="margin-left:auto">P to exit</span>
  </div>
</aside>
<script>
(function () {
  var slides = Array.prototype.slice.call(document.querySelectorAll(".slide"));
  var cur = 0, notesOn = false;
  var progressBar = document.getElementById("progress-bar");
  var curEl = document.getElementById("cur");
  var toast = document.getElementById("toast");
  document.getElementById("total").textContent = slides.length;
  function show(i) {
    cur = Math.max(0, Math.min(slides.length - 1, i));
    slides.forEach(function (s, k) { s.classList.toggle("active", k === cur); });
    progressBar.style.width = ((cur + 1) / slides.length * 100) + "%";
    curEl.textContent = cur + 1;
    if (presenterOn) renderPresenter();
    else if (notesOn) renderNotes();
    else toast.classList.remove("show");
    history.replaceState(null, "", "#" + (cur + 1));
  }
  var NOTES = __NOTES__;
  function renderNotes() {
    var text = NOTES[cur] || "";
    if (!text) { toast.classList.remove("show"); return; }
    toast.textContent = text;
    toast.classList.add("show");
  }
  function toggleNotes() {
    notesOn = !notesOn;
    document.getElementById("notes").classList.toggle("on", notesOn);
    if (notesOn) renderNotes(); else toast.classList.remove("show");
  }
  var presenterOn = false, timerStart = null, timerInt = null;
  function currentCardSize() {
    var w = slides[cur].getBoundingClientRect().width;
    return { w: w, h: w * 9 / 16 };
  }
  function renderPresenter() {
    if (!presenterOn) return;
    var stage = document.getElementById("pv-next-stage");
    var holder = document.getElementById("pv-next-holder");
    holder.innerHTML = "";
    var next = slides[cur + 1];
    var size = currentCardSize();
    if (!next) {
      holder.innerHTML = '<div class="pv-end">End of deck \u2014 Q&amp;A</div>';
    } else {
      var clone = next.cloneNode(true);
      clone.classList.add("active");
      clone.style.position = "absolute";
      clone.style.margin = "0";
      clone.style.width = size.w + "px";
      clone.style.height = size.h + "px";
      var k = stage.clientWidth / size.w;
      clone.style.transform = "scale(" + k + ")";
      clone.style.transformOrigin = "top left";
      holder.appendChild(clone);
    }
    var title = slides[cur].querySelector(".slide-title");
    document.getElementById("pv-title").textContent = title ? title.textContent : "Slide " + (cur + 1);
    document.getElementById("pv-body").textContent = NOTES[cur] || "(no notes for this slide)";
    document.getElementById("pv-count").textContent = (cur + 1) + " / " + slides.length;
  }
  function tickTimer() {
    var s = Math.floor((Date.now() - timerStart) / 1000);
    var m = Math.floor(s / 60); s = s % 60;
    document.getElementById("pv-timer").textContent = (m < 10 ? "0" : "") + m + ":" + (s < 10 ? "0" : "") + s;
  }
  function togglePresenter(force) {
    presenterOn = typeof force === "boolean" ? force : !presenterOn;
    document.body.classList.toggle("presenter", presenterOn);
    document.getElementById("present").classList.toggle("on", presenterOn);
    if (presenterOn) {
      toast.classList.remove("show");
      timerStart = Date.now();
      tickTimer();
      if (timerInt) clearInterval(timerInt);
      timerInt = setInterval(tickTimer, 1000);
      renderPresenter();
    } else if (timerInt) {
      clearInterval(timerInt);
      timerInt = null;
    }
  }
  document.getElementById("prev").onclick = function () { show(cur - 1); };
  document.getElementById("next").onclick = function () { show(cur + 1); };
  document.getElementById("notes").onclick = toggleNotes;
  document.getElementById("present").onclick = function () { togglePresenter(); };
  window.addEventListener("resize", function () { renderPresenter(); });
  document.getElementById("help").onclick = function () {
    alert("Navigate: \u2190 \u2192 arrows, Space, PgUp/PgDn\nFirst / last: Home / End\nN: toggle speaker notes toast\nP: presenter view \u2014 current + next slide + notes + timer\nURL hash: open at any slide, e.g. #7\nPrint: browser Print to export all slides as PDF");
  };
  document.addEventListener("keydown", function (e) {
    if (e.key === "ArrowRight" || e.key === "PageDown" || e.key === " ") { show(cur + 1); e.preventDefault(); }
    else if (e.key === "ArrowLeft" || e.key === "PageUp") { show(cur - 1); e.preventDefault(); }
    else if (e.key === "Home") { show(0); e.preventDefault(); }
    else if (e.key === "End") { show(slides.length - 1); e.preventDefault(); }
    else if (e.key === "n" || e.key === "N") { toggleNotes(); }
    else if (e.key === "p" || e.key === "P") { togglePresenter(); }
  });
  function fromHash() {
    var n = parseInt(location.hash.slice(1), 10);
    show(isNaN(n) ? 0 : n - 1);
  }
  window.addEventListener("hashchange", fromHash);
  fromHash();
})();
</script>
</body>
</html>
"""


def cover_meta(slides):
    """(tagline, meta lines) for the title slide, shared by the PPTX and HTML
    builders so the two covers cannot drift apart."""
    bullets = slides[0]["bullets"]
    tagline = ""
    for b in bullets:
        if b.startswith("Tagline:"):
            tagline = strip_md(b[len("Tagline:"):].strip())
            break
        if b.startswith("From donor"):
            tagline = strip_md(b)
            break
    meta = [
        strip_md(b)
        for b in bullets
        if not b.startswith("Tagline:") and not b.startswith("From donor") and not b.startswith("CBODS —")
    ]
    return tagline, meta


def build_html_deck(slides):
    font_path = ROOT / "static" / "fonts" / "InterVariable.woff2"
    font_b64 = base64.b64encode(font_path.read_bytes()).decode("ascii")

    parts = []
    for data in slides:
        if data["number"] == 1:
            tagline, meta_lines = cover_meta(slides)
            meta_lines = [m for m in meta_lines if m]
            meta_text = "\n".join(md_inline(m) for m in meta_lines)
            parts.append(
                '<section class="slide cover" aria-label="Title">\n'
                '  <h1 class="slide-title">CBODS</h1>\n'
                '  <p class="sub">Centralised Blood &amp; Organ Donation System</p>\n'
                '  <hr class="rule">\n'
                f'  <p class="tagline">{md_inline(tagline)}</p>\n'
                f'  <p class="meta">{meta_text}</p>\n'
                "</section>"
            )
            continue
        bullets = "\n".join(
            f'  <li><span class="marker" aria-hidden="true">•</span><span>{md_inline(b)}</span></li>'
            for b in data["bullets"]
        )
        parts.append(
            f'<section class="slide" aria-label="Slide {data["number"]}: {strip_md(data["title"])}">\n'
            f'  <span class="chip">{data["time"]}</span>\n'
            f'  <h2 class="slide-title">{md_inline(data["title"])}</h2>\n'
            '  <hr class="rule">\n'
            '  <ul class="bullets">\n'
            f"{bullets}\n"
            "  </ul>\n"
            '  <footer class="slide-foot">\n'
            '    <span>CBODS — Centralised Blood &amp; Organ Donation System</span>\n'
            f'    <span>{data["number"]} / {len(slides)}</span>\n'
            "  </footer>\n"
            "</section>"
        )
    body = "\n".join(parts)
    notes_json = json.dumps([strip_md(s["notes"]) for s in slides], ensure_ascii=False)
    html = (
        HTML_TEMPLATE
        .replace("__INTER_B64__", font_b64)
        .replace("__SLIDES__", body)
        .replace("__TOTAL__", str(len(slides)))
        .replace("__NOTES__", notes_json)
    )
    out = DOCS / "CBODS_PRESENTATION.html"
    out.write_text(html, encoding="utf-8")
    return out, len(slides)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def build_pdf():
    from fpdf import FPDF

    md = (
        (DOCS / "PANEL_PREP.md")
        .read_text(encoding="utf-8")
        .replace("↔", "<->")
        .replace("⇒", "=>")
    )

    BRAND = (164, 22, 26)
    DARK = (34, 34, 34)
    GREY = (90, 90, 90)
    LIGHT = (243, 222, 225)

    fonts = {}
    for style, filename in [("", "arial.ttf"), ("B", "arialbd.ttf"), ("I", "ariali.ttf"), ("BI", "arialbi.ttf")]:
        path = pathlib.Path(r"C:\Windows\Fonts") / filename
        if path.exists():
            fonts[style] = str(path)
    if not fonts:
        sys.exit("Arial fonts not found in C:\\Windows\\Fonts")

    class PanelPDF(FPDF):
        def footer(self):
            self.set_y(-15)
            self.set_font("CBODS", "I", 8.5)
            self.set_text_color(*GREY)
            self.cell(0, 10, f"CBODS — Panel Prep · Page {self.page_no()}/{{nb}}", align="C")

    pdf = PanelPDF(format="A4")
    pdf.alias_nb_pages()
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_font("CBODS", "", fonts[""])
    if "B" in fonts:
        pdf.add_font("CBODS", "B", fonts["B"])
    if "I" in fonts:
        pdf.add_font("CBODS", "I", fonts["I"])
    if "BI" in fonts:
        pdf.add_font("CBODS", "BI", fonts["BI"])

    def h1(text):
        pdf.set_font("CBODS", "B", 21)
        pdf.set_text_color(*BRAND)
        pdf.multi_cell(0, 11, text)
        pdf.set_draw_color(*BRAND)
        pdf.set_line_width(0.6)
        pdf.line(18, pdf.get_y() + 2, 192, pdf.get_y() + 2)
        pdf.ln(7)

    def h2(text):
        if pdf.get_y() > 240:
            pdf.add_page()
        pdf.ln(2)
        pdf.set_font("CBODS", "B", 13.5)
        pdf.set_text_color(*BRAND)
        pdf.multi_cell(0, 8, text)
        pdf.set_draw_color(*LIGHT)
        pdf.set_line_width(0.5)
        pdf.line(18, pdf.get_y() + 1, 192, pdf.get_y() + 1)
        pdf.ln(4)

    def para(text, italic=False, grey=False):
        pdf.set_font("CBODS", "BI" if italic and "BI" in fonts else ("I" if italic else ""), 10.5)
        pdf.set_text_color(*(GREY if grey else DARK))
        pdf.multi_cell(0, 6, text)
        pdf.ln(1.5)

    def quote(text):
        y = pdf.get_y()
        pdf.set_fill_color(*LIGHT)
        pdf.set_font("CBODS", "I", 10.5)
        pdf.set_text_color(*DARK)
        pdf.set_x(24)
        pdf.multi_cell(160, 6, text, fill=True)
        pdf.set_fill_color(*BRAND)
        pdf.rect(19.5, y, 1.6, pdf.get_y() - y, "F")
        pdf.ln(2)

    def bullets(items):
        for item in items:
            if pdf.get_y() > 262:
                pdf.add_page()
            pdf.set_font("CBODS", "", 10.5)
            pdf.set_x(24)
            pdf.set_text_color(*BRAND)
            pdf.cell(5, 6, "•")
            pdf.set_x(29)
            pdf.set_text_color(*DARK)
            pdf.multi_cell(157, 6, item, markdown=True)
            pdf.ln(0.8)

    def table(header, rows):
        col1_w, col2_w, row_h = 52, 140, 7
        if pdf.get_y() > 250:
            pdf.add_page()
        pdf.set_draw_color(*LIGHT)
        pdf.set_line_width(0.2)
        pdf.set_font("CBODS", "B", 10)
        pdf.set_text_color(255, 255, 255)
        pdf.set_fill_color(*BRAND)
        pdf.cell(col1_w, row_h, header[0], border=1, fill=True)
        pdf.cell(col2_w, row_h, header[1], border=1, fill=True)
        pdf.ln(row_h)
        for i, row in enumerate(rows):
            if pdf.get_y() > 262:
                pdf.add_page()
            fill = i % 2 == 0
            pdf.set_fill_color(248, 244, 245)
            pdf.set_text_color(*DARK)
            pdf.set_font("CBODS", "B", 10)
            pdf.cell(col1_w, row_h, strip_md(row[0]), border=1, fill=fill)
            pdf.set_font("CBODS", "", 10)
            pdf.cell(col2_w, row_h, strip_md(row[1]), border=1, fill=fill)
            pdf.ln(row_h)
        pdf.ln(2)

    pdf.add_page()
    lines = md.splitlines()
    title = lines[0].lstrip("# ").strip()
    intro = [ln.strip() for ln in lines[1:] if ln.strip() and not ln.startswith("#")]
    h1(title)
    for ln in intro[:3]:
        para(ln, italic=True, grey=True)

    sections = re.split(r"^## ", md, flags=re.M)[1:]
    for section in sections:
        sec_lines = section.splitlines()
        heading = sec_lines[0].strip()
        if heading.lower().startswith("cbods — panel prep"):
            continue
        h2(heading)
        bullets_buf, table_buf = [], []
        table_mode = False

        def flush_table():
            nonlocal table_buf
            if table_buf:
                header = [c.strip().strip("*") for c in table_buf[0].strip().strip("|").split("|")]
                rows = []
                for line in table_buf[1:]:
                    cells = [c.strip() for c in line.strip().strip("|").split("|")]
                    if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                        continue
                    rows.append(cells)
                table(header, rows)
                table_buf = []

        for raw in sec_lines[1:]:
            line = raw.rstrip()
            stripped = line.strip()
            if stripped.startswith("|"):
                table_mode = True
                table_buf.append(stripped)
                continue
            if table_mode and not stripped.startswith("|"):
                flush_table()
                table_mode = False
            if stripped.startswith("- "):
                bullets_buf.append(stripped[2:])
            elif stripped.startswith("> "):
                if bullets_buf:
                    bullets(bullets_buf)
                    bullets_buf = []
                quote(strip_md(stripped[2:]))
            elif stripped:
                if bullets_buf:
                    bullets(bullets_buf)
                    bullets_buf = []
                para(stripped)
        if table_mode:
            flush_table()
        if bullets_buf:
            bullets(bullets_buf)

    out = DOCS / "PANEL_PREP.pdf"
    pdf.output(out)
    return out, pdf.page_no()


if __name__ == "__main__":
    slides = parse_outline()
    assert len(slides) == 10, f"expected 10 slides, found {len(slides)}"
    html_path, count = build_html_deck(slides)
    pptx_path, count = build_pptx(slides)
    pdf_path, pages = build_pdf()
    print(f"OK {html_path.relative_to(ROOT)} ({count} slides, self-contained)")
    print(f"OK {pptx_path.relative_to(ROOT)} ({count} slides, notes on every slide)")
    print(f"OK {pdf_path.relative_to(ROOT)} ({pages} pages)")
